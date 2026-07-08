"""Bob (Mesa) — placement intelligence workspace API.

Shared-workspace auth: one access code per deployment (BOB_ACCESS_CODE),
sent as X-Bob-Key on every request. No per-user accounts in v1.
"""

import io
import json
import logging

from fastapi import APIRouter, Depends, File, Header, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from sqlalchemy import text
from sqlalchemy.orm import Session

from core.config import settings
from database.session import get_db
from services.bob.schema import ensure_schema
from services.bob.agent import start_run

logger = logging.getLogger(__name__)
router = APIRouter(prefix="/bob", tags=["Bob"])


# ── Auth + schema bootstrap ────────────────────────────────────────────────────

def _guard(x_bob_key: str = Header(default="")) -> None:
    if not settings.BOB_ACCESS_CODE:
        raise HTTPException(status_code=503, detail="Bob is not configured on this environment")
    if x_bob_key != settings.BOB_ACCESS_CODE:
        raise HTTPException(status_code=401, detail="Invalid access code")
    ensure_schema()


class VerifyRequest(BaseModel):
    code: str


@router.post("/auth/verify")
async def verify_code(req: VerifyRequest):
    if not settings.BOB_ACCESS_CODE:
        raise HTTPException(status_code=503, detail="Bob is not configured on this environment")
    if req.code != settings.BOB_ACCESS_CODE:
        raise HTTPException(status_code=401, detail="Invalid access code")
    ensure_schema()
    return {"ok": True}


# ── Chats ──────────────────────────────────────────────────────────────────────

@router.get("/chats", dependencies=[Depends(_guard)])
async def list_chats(db: Session = Depends(get_db)):
    rows = db.execute(
        text("SELECT id, title, created_at, updated_at FROM bob_chats ORDER BY updated_at DESC LIMIT 200")
    ).fetchall()
    return {"chats": [
        {"id": r[0], "title": r[1], "created_at": str(r[2]), "updated_at": str(r[3])} for r in rows
    ]}


@router.post("/chats", dependencies=[Depends(_guard)])
async def create_chat(db: Session = Depends(get_db)):
    row = db.execute(text("INSERT INTO bob_chats (title) VALUES ('New chat') RETURNING id")).fetchone()
    db.commit()
    return {"id": row[0], "title": "New chat"}


@router.delete("/chats/{chat_id}", dependencies=[Depends(_guard)])
async def delete_chat(chat_id: int, db: Session = Depends(get_db)):
    db.execute(text("DELETE FROM bob_chats WHERE id = :c"), {"c": chat_id})
    db.commit()
    return {"ok": True}


@router.get("/chats/{chat_id}", dependencies=[Depends(_guard)])
async def get_chat(chat_id: int, db: Session = Depends(get_db)):
    chat = db.execute(
        text("SELECT id, title FROM bob_chats WHERE id = :c"), {"c": chat_id}
    ).fetchone()
    if not chat:
        raise HTTPException(status_code=404, detail="Chat not found")

    messages = db.execute(
        text("SELECT id, role, content, created_at, meta FROM bob_messages WHERE chat_id = :c ORDER BY id"),
        {"c": chat_id},
    ).fetchall()
    run = db.execute(
        text("SELECT id, status, events, counters, credits_used, answer FROM bob_runs "
             "WHERE chat_id = :c ORDER BY id DESC LIMIT 1"),
        {"c": chat_id},
    ).fetchone()

    return {
        "id": chat[0],
        "title": chat[1],
        "messages": [
            {"id": m[0], "role": m[1], "content": m[2], "created_at": str(m[3]), "meta": m[4] or {}}
            for m in messages
        ],
        "latest_run": _run_payload(run) if run else None,
        "tables": _tables_payload(db, chat_id),
    }


# ── File attachments ───────────────────────────────────────────────────────────

MAX_UPLOAD_BYTES = 10 * 1024 * 1024  # 10 MB


@router.post("/chats/{chat_id}/files", dependencies=[Depends(_guard)])
async def upload_chat_file(chat_id: int, file: UploadFile = File(...), db: Session = Depends(get_db)):
    from services.bob.files import extract_text, FileExtractionError

    chat = db.execute(text("SELECT id FROM bob_chats WHERE id = :c"), {"c": chat_id}).fetchone()
    if not chat:
        raise HTTPException(status_code=404, detail="Chat not found")

    data = await file.read()
    if len(data) > MAX_UPLOAD_BYTES:
        raise HTTPException(status_code=413, detail="File too large (max 10 MB)")

    try:
        extracted = extract_text(file.filename or "upload", data)
    except FileExtractionError as e:
        raise HTTPException(status_code=422, detail=str(e))

    row = db.execute(
        text("INSERT INTO bob_files (chat_id, filename, mime, text_content) "
             "VALUES (:c, :f, :m, :t) RETURNING id"),
        {"c": chat_id, "f": (file.filename or "upload")[:200],
         "m": (file.content_type or "")[:100], "t": extracted},
    ).fetchone()
    db.commit()
    return {"file_id": row[0], "filename": file.filename, "chars_extracted": len(extracted)}


# ── Messages / runs ────────────────────────────────────────────────────────────

class SendMessageRequest(BaseModel):
    content: str


@router.post("/chats/{chat_id}/messages", dependencies=[Depends(_guard)])
async def send_message(chat_id: int, req: SendMessageRequest, db: Session = Depends(get_db)):
    content = (req.content or "").strip()
    if not content:
        raise HTTPException(status_code=400, detail="Empty message")
    if len(content) > 8000:
        content = content[:8000]

    chat = db.execute(text("SELECT id, title FROM bob_chats WHERE id = :c"), {"c": chat_id}).fetchone()
    if not chat:
        raise HTTPException(status_code=404, detail="Chat not found")

    # A run whose thread died (pod restart, crash) stays 'running' forever —
    # treat anything silent for 15+ minutes as dead so the chat never wedges.
    running = db.execute(
        text("SELECT id FROM bob_runs WHERE chat_id = :c AND status = 'running' "
             "AND updated_at > now() - interval '15 minutes'"),
        {"c": chat_id},
    ).fetchone()
    if running:
        raise HTTPException(status_code=409, detail="Bob is still working on this chat")
    db.execute(
        text("UPDATE bob_runs SET status = 'error', error = 'stale run reaped', updated_at = now() "
             "WHERE chat_id = :c AND status = 'running'"),
        {"c": chat_id},
    )

    db.execute(
        text("INSERT INTO bob_messages (chat_id, role, content) VALUES (:c, 'user', :m)"),
        {"c": chat_id, "m": content},
    )
    if chat[1] == "New chat":
        db.execute(
            text("UPDATE bob_chats SET title = :t, updated_at = now() WHERE id = :c"),
            {"t": content[:70], "c": chat_id},
        )
    else:
        db.execute(text("UPDATE bob_chats SET updated_at = now() WHERE id = :c"), {"c": chat_id})
    db.commit()

    run_id = start_run(chat_id)
    return {"run_id": run_id}


@router.get("/runs/{run_id}", dependencies=[Depends(_guard)])
async def get_run(run_id: int, db: Session = Depends(get_db)):
    run = db.execute(
        text("SELECT id, status, events, counters, credits_used, answer, chat_id FROM bob_runs WHERE id = :r"),
        {"r": run_id},
    ).fetchone()
    if not run:
        raise HTTPException(status_code=404, detail="Run not found")
    payload = _run_payload(run)
    payload["tables"] = _tables_payload(db, run[6])
    return payload


# ── Rows: status updates (team coordination) ──────────────────────────────────

class RowStatusRequest(BaseModel):
    status: str


@router.patch("/rows/{row_id}", dependencies=[Depends(_guard)])
async def update_row_status(row_id: int, req: RowStatusRequest, db: Session = Depends(get_db)):
    if req.status not in ("new", "contacted", "replied", "meeting", "dead"):
        raise HTTPException(status_code=400, detail="Invalid status")
    db.execute(
        text("UPDATE bob_rows SET status = :s, updated_at = now() WHERE id = :r"),
        {"s": req.status, "r": row_id},
    )
    db.commit()
    return {"ok": True}


@router.delete("/rows/{row_id}", dependencies=[Depends(_guard)])
async def delete_row(row_id: int, db: Session = Depends(get_db)):
    # Record the removal in the rejection ledger BEFORE deleting: a company the
    # user removed must never be re-added by a later agent run.
    from services.bob.agent import _norm_company

    meta = db.execute(
        text("SELECT r.cells->>'company', t.chat_id FROM bob_rows r "
             "JOIN bob_tables t ON t.id = r.table_id WHERE r.id = :r"),
        {"r": row_id},
    ).fetchone()
    if meta and meta[0]:
        db.execute(
            text("INSERT INTO bob_rejections (chat_id, company, company_norm, reason) "
                 "VALUES (:c, :co, :n, 'removed by user')"),
            {"c": meta[1], "co": meta[0], "n": _norm_company(meta[0])},
        )
    res = db.execute(text("DELETE FROM bob_rows WHERE id = :r"), {"r": row_id})
    db.commit()
    if res.rowcount == 0:
        raise HTTPException(status_code=404, detail="Row not found")
    return {"ok": True, "deleted": row_id, "banned": bool(meta and meta[0])}


class AddColumnsRequest(BaseModel):
    columns: list[dict]  # [{"key": "...", "label": "..."}]


@router.post("/tables/{table_id}/columns", dependencies=[Depends(_guard)])
async def add_table_columns(table_id: int, req: AddColumnsRequest, db: Session = Depends(get_db)):
    """Append columns to a table (admin/import use). Merges by key, no dupes."""
    existing = db.execute(text("SELECT columns FROM bob_tables WHERE id=:t"), {"t": table_id}).scalar()
    if existing is None:
        raise HTTPException(status_code=404, detail="Table not found")
    keys = {c.get("key") for c in existing}
    merged = list(existing) + [c for c in req.columns if c.get("key") and c.get("key") not in keys]
    db.execute(text("UPDATE bob_tables SET columns=CAST(:c AS jsonb), updated_at=now() WHERE id=:t"),
               {"c": json.dumps(merged), "t": table_id})
    db.commit()
    return {"ok": True, "columns": [c.get("key") for c in merged]}


class AddRowsRequest(BaseModel):
    rows: list[dict]  # each keyed by column key


@router.post("/tables/{table_id}/rows", dependencies=[Depends(_guard)])
async def add_table_rows(table_id: int, req: AddRowsRequest, db: Session = Depends(get_db)):
    """Append rows directly (admin/import use — bypasses the agent). Cells are
    sanitized (URL validation, em-dash scrub) like the agent path."""
    from services.bob.agent import _sanitize_cells

    if db.execute(text("SELECT 1 FROM bob_tables WHERE id=:t"), {"t": table_id}).scalar() is None:
        raise HTTPException(status_code=404, detail="Table not found")
    pos = db.execute(text("SELECT coalesce(max(position),0) FROM bob_rows WHERE table_id=:t"),
                     {"t": table_id}).scalar()
    added = []
    for i, r in enumerate(req.rows):
        if not isinstance(r, dict):
            continue
        clean, _ = _sanitize_cells(r)
        row = db.execute(
            text("INSERT INTO bob_rows (table_id, position, cells) "
                 "VALUES (:t, :p, CAST(:c AS jsonb)) RETURNING id"),
            {"t": table_id, "p": pos + i + 1, "c": json.dumps(clean, ensure_ascii=False)},
        ).fetchone()
        added.append(row[0])
    db.execute(text("UPDATE bob_tables SET updated_at=now() WHERE id=:t"), {"t": table_id})
    db.commit()
    return {"ok": True, "added": added, "count": len(added)}


class GuardrailsRequest(BaseModel):
    reject_companies: list[str] = []
    target_functions: list[str] = []


@router.post("/chats/{chat_id}/guardrails", dependencies=[Depends(_guard)])
async def set_guardrails(chat_id: int, req: GuardrailsRequest, db: Session = Depends(get_db)):
    """Seed the rejection ledger and/or the mandate function keywords for a chat."""
    from services.bob.agent import _norm_company

    added = 0
    for c in req.reject_companies:
        c = (c or "").strip()
        if not c:
            continue
        n = _norm_company(c)
        exists = db.execute(
            text("SELECT 1 FROM bob_rejections WHERE chat_id = :c AND company_norm = :n LIMIT 1"),
            {"c": chat_id, "n": n},
        ).fetchone()
        if not exists:
            db.execute(
                text("INSERT INTO bob_rejections (chat_id, company, company_norm, reason) "
                     "VALUES (:c, :co, :n, 'removed by user')"),
                {"c": chat_id, "co": c, "n": n},
            )
            added += 1
    tf = [t.strip().lower() for t in req.target_functions if t.strip()][:12]
    tables = 0
    if tf:
        res = db.execute(
            text("UPDATE bob_tables SET mandate = CAST(:m AS jsonb) WHERE chat_id = :c"),
            {"m": json.dumps(tf), "c": chat_id},
        )
        tables = res.rowcount
    db.commit()
    return {"ok": True, "rejections_added": added, "tables_mandated": tables}


class RowCellsRequest(BaseModel):
    cells: dict


@router.patch("/rows/{row_id}/cells", dependencies=[Depends(_guard)])
async def update_row_cells(row_id: int, req: RowCellsRequest, db: Session = Depends(get_db)):
    """Direct cell repair (sanitized). Used for audit fixes; the agent has its own tool."""
    from services.bob.agent import _sanitize_cells

    clean, removed = _sanitize_cells(dict(req.cells or {}))
    res = db.execute(
        text("UPDATE bob_rows SET cells = cells || CAST(:c AS jsonb), updated_at = now() WHERE id = :r"),
        {"c": json.dumps(clean, ensure_ascii=False), "r": row_id},
    )
    db.commit()
    if res.rowcount == 0:
        raise HTTPException(status_code=404, detail="Row not found")
    return {"ok": True, "row_id": row_id, "removed_by_validation": removed}


@router.delete("/tables/{table_id}", dependencies=[Depends(_guard)])
async def delete_table(table_id: int, db: Session = Depends(get_db)):
    db.execute(text("DELETE FROM bob_rows WHERE table_id = :t"), {"t": table_id})
    res = db.execute(text("DELETE FROM bob_tables WHERE id = :t"), {"t": table_id})
    db.commit()
    if res.rowcount == 0:
        raise HTTPException(status_code=404, detail="Table not found")
    return {"ok": True, "deleted": table_id}


@router.post("/tables/{table_id}/scrub", dependencies=[Depends(_guard)])
async def scrub_table(table_id: int, db: Session = Depends(get_db)):
    """Re-run the cell sanitizer over every row (applies rails added after the
    rows were written, e.g. X links in linkedin columns)."""
    from services.bob.agent import _sanitize_cells

    rows = db.execute(
        text("SELECT id, cells FROM bob_rows WHERE table_id = :t"), {"t": table_id}
    ).fetchall()
    fixed = []
    for rid, cells in rows:
        clean, removed = _sanitize_cells(dict(cells or {}))
        if removed:
            db.execute(
                text("UPDATE bob_rows SET cells = CAST(:c AS jsonb), updated_at = now() WHERE id = :r"),
                {"c": json.dumps(clean, ensure_ascii=False), "r": rid},
            )
            fixed.append({"row_id": rid, "removed": removed})
    db.commit()
    return {"ok": True, "rows_fixed": len(fixed), "details": fixed}


# ── Export ─────────────────────────────────────────────────────────────────────

@router.get("/tables/{table_id}/export", dependencies=[Depends(_guard)])
async def export_table(table_id: int, db: Session = Depends(get_db)):
    tbl = db.execute(
        text("SELECT name, columns FROM bob_tables WHERE id = :t"), {"t": table_id}
    ).fetchone()
    if not tbl:
        raise HTTPException(status_code=404, detail="Table not found")
    rows = db.execute(
        text("SELECT cells, status FROM bob_rows WHERE table_id = :t ORDER BY position, id"),
        {"t": table_id},
    ).fetchall()

    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = (tbl[0] or "Results")[:31]
    columns = tbl[1] or []
    headers = [c.get("label") or c.get("key") for c in columns] + ["Status"]
    ws.append(headers)
    for cells, status in rows:
        cells = cells or {}
        ws.append([_cell_str(cells.get(c.get("key"))) for c in columns] + [status])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    fname = (tbl[0] or "results").replace('"', "")[:60]
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{fname}.xlsx"'},
    )


# ── Helpers ────────────────────────────────────────────────────────────────────

def _cell_str(v) -> str:
    if v is None:
        return ""
    if isinstance(v, (dict, list)):
        return json.dumps(v, ensure_ascii=False)
    return str(v)


def _run_payload(run) -> dict:
    return {
        "id": run[0],
        "status": run[1],
        "events": run[2] or [],
        "counters": run[3] or {},
        "credits_used": run[4] or 0,
        "answer": run[5] or "",
    }


def _tables_payload(db: Session, chat_id: int) -> list[dict]:
    tables = db.execute(
        text("SELECT id, name, columns FROM bob_tables WHERE chat_id = :c ORDER BY id"),
        {"c": chat_id},
    ).fetchall()
    out = []
    for tid, name, columns in tables:
        rows = db.execute(
            text("SELECT id, cells, status FROM bob_rows WHERE table_id = :t ORDER BY position, id LIMIT 500"),
            {"t": tid},
        ).fetchall()
        out.append({
            "id": tid,
            "name": name,
            "columns": columns or [],
            "rows": [{"id": r[0], "cells": r[1] or {}, "status": r[2]} for r in rows],
        })
    return out
